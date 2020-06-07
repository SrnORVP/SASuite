
import docx, os, SACusFun as SACF, pandas as pd, numpy as np, main_pandas
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from datetime import datetime as dt

arrPath = main_pandas.arrPath
strSACInput = main_pandas.strSACinput
strSASparam = main_pandas.strSASparam
strSACOutput = main_pandas.strSACoutput


# Docx Blocks Generator
def iter_Docx_Block_Items(parent):
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent. _tc
    else:
        raise ValueError("Something's not right")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


# Extract Word Docx
def extract_Style_Index_From_Docs(docxSIF):
    arrBlkStyle = []
    arrObjBlk = []
    seqTblBlk = 0
    arrSAC = []
    strSubsysGrp = ''
    for objBlk in iter_Docx_Block_Items(docxSIF):
        if isinstance(objBlk, Paragraph):
            strBlkStyle = objBlk.style.name
            if strBlkStyle == 'Heading 1':
                strBlkSIF = objBlk.text
            if strBlkStyle == 'Heading 3' or strBlkStyle == 'Heading 4':
                strSubsysGrp = objBlk.text.strip().replace("  ", " ")
        else:
            seqTblBlk += 1
            isOverview = 'Equipment Leg' in objBlk.cell(0,0).text
            isFRdata = 'Component' in objBlk.cell(0,0).text
            if isOverview:
                strTblType = 'Overview'
            elif isFRdata:
                strTblType = 'FRdata'
            if isOverview or isFRdata:
                for seqRow, objRow in enumerate(objBlk.rows):
                    arrSACRow = []
                    arrSACRow.append(strBlkSIF)
                    arrSACRow.append(seqTblBlk)
                    arrSACRow.append(strSubsysGrp)
                    arrSACRow.append(strTblType)
                    arrSACRow.append(seqRow)
                    for objCell in objRow.cells:
                        arrSACRow.append(objCell.text)
                    arrSAC.append(arrSACRow)
    ndaSAC = np.array(arrSAC)
    return ndaSAC


# Extract String
def extract_String_on_Conditions(strInput,arrCond):
    for elemCond in arrCond:
        strExtracted = SACF.extract_Substring(strInput,elemCond[0],elemCond[1])
        if strInput != strExtracted:
             return strExtracted
    return strInput


# Load Inputs
docxSIF = docx.Document(os.path.join(*arrPath, strSACInput))
strSASIP = os.path.join('.', strSASparam)
arrC4SAC_str = pd.read_excel(strSASIP,'C4SAC-str').values
arrC4SAC_tbl = pd.read_excel(strSASIP,'C4SAC-tbl').values
arrC4SAC_hdr = pd.read_excel(strSASIP,'C4SAC-hdr').values


# Extract Word Doc
ndaSAC = extract_Style_Index_From_Docs(docxSIF)
dfSAC = pd.DataFrame(ndaSAC)


# Expand List
dfSAC = dfSAC[0].apply(pd.Series)


# CleanUp - str
arrExStr1 = arrC4SAC_str[arrC4SAC_str[:,2]==1]
arrExStr2 = arrC4SAC_str[arrC4SAC_str[:,2]==2]

dfSAC['Group'] = dfSAC[2].apply(extract_String_on_Conditions,1,arrCond=arrExStr1)
dfSAC[2] = dfSAC[2].apply(extract_String_on_Conditions,1,arrCond=arrExStr2)


# CleanUp - Column Labels
dfSAC = dfSAC.rename({0:'SIF',1:'Table',2:'Tag',3:'Type',4:'Row',5:1,6:2,7:3,8:4,9:5,10:6,11:7,12:8,13:9,14:10,15:11},axis=1)
dfSAC = dfSAC.reindex(['SIF','Group','Tag','Table','Type','Row',1,2,3,4,5,6,7,8,9,10,11],axis=1)


# Separate into different 5 table types, and remove head and tail rows for each table based on `arrC4SAC_tbl`
dictSAC = dict()
for rowC4SAC_tbl in arrC4SAC_tbl:
    dfSAC_tbl = pd.DataFrame()
    dftemp = dfSAC[(dfSAC['Group'] == rowC4SAC_tbl[1]) & (dfSAC['Type'] == rowC4SAC_tbl[2])]
    arrtemp_enumTbl = dftemp['Table'].unique()
    dftemp_grped = dftemp.groupby('Table')
    for enumTbl in arrtemp_enumTbl:
        if rowC4SAC_tbl[4] == 0:
            intEnd = None
        else:
            intEnd = -rowC4SAC_tbl[4]
        dfSAC_tbl = dfSAC_tbl.append(dftemp_grped.get_group(enumTbl).iloc[rowC4SAC_tbl[3]:intEnd])
    dictSAC[rowC4SAC_tbl[0]] = dfSAC_tbl


# Renaming Columns based on `arrC4SAC_hdr`
for rowC4SAC_hdr in arrC4SAC_hdr:
    dictSAC_hdr = dict()
    for seqC4SAC_hdr in range(1,len(rowC4SAC_hdr)):
        dictSAC_hdr[seqC4SAC_hdr] = rowC4SAC_hdr[seqC4SAC_hdr]
    dictSAC[rowC4SAC_hdr[0]] = dictSAC[rowC4SAC_hdr[0]].rename(dictSAC_hdr, axis=1)


# Remove PLC DataFrame from other DataFrame
# It is treated differently from now on
dfSAC_PLC = dictSAC.pop('PLC_FRD')


# Remove Unnecessary Columns
dictSAC_Trim = dict()
for strkey_dictSAC in dictSAC.keys():
    dfTemp = dictSAC[strkey_dictSAC].drop(['Group','Type','Row'],axis=1).dropna(axis=1,how='all')
    dictSAC_Trim[strkey_dictSAC] = dfTemp.set_index('Tag')


# Join DataFrame for each subsystems
dictSAC_Trim['INT_FRD'] = dictSAC_Trim['INT_FRD'].drop(['SIF','Table'],axis=1)
dictSAC_Trim['PFE_FRD'] = dictSAC_Trim['PFE_FRD'].drop(['SIF','Table'],axis=1)
dfSAC_INT = pd.concat([dictSAC_Trim['INT_ORV'], dictSAC_Trim['INT_FRD']],axis=1,sort=False)
dfSAC_PFE = pd.concat([dictSAC_Trim['PFE_ORV'], dictSAC_Trim['PFE_FRD']],axis=1,sort=False)

dfSAC_INT = dfSAC_INT.sort_values(['Tag','Table'], axis=0)
dfSAC_PFE = dfSAC_PFE.sort_values(['Tag','Table'], axis=0)


# Export SAC
strTimeNow = dt.now().strftime("%d%b%y")
with pd.ExcelWriter(f'{os.path.join(*arrPath, "")}{strSACOutput}-{strTimeNow}.xlsx') as writer:
    dfSAC_INT.to_excel(writer,'INT')
    dfSAC_PFE.to_excel(writer,'PFE')
    dfSAC_PLC.to_excel(writer,'PLC')

